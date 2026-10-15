'Dapur' os
import shutil
import glob
import subprocess
from openpyxl import load_workbook

def main():
    source_file = 'Form Revisi.xlsx'
    dapur_folder = 'Dapur'
    addon_folder ='Addon'
    
    print('--> Mengecek keberadaan file sumber...')
    if not os.path.exists(source_file):
        print('--> Gagal: File "Form Revisi.xlsx" tidak ditemukan.')
        input('--> Tekan enter untuk keluar')
        return

    try:
        wb = load_workbook(source_file, read_only=True)
        if len(wb.sheetnames) > 1:
            print('--> File tersebut memiliki banyak sheet, hapus dan sisakan yang akan di olah.')
            input('--> Tekan enter untuk keluar')
            return
        wb.close()
    except Exception as e:
        print(f'--> Gagal membaca file excel: {e}')
        input('--> Tekan enter untuk keluar')
        return

    print('--> Mengecek folder dan dependensi Dapur...')
    if not os.path.exists(dapur_folder):
        print('--> Gagal: Folder "Dapur" tidak ditemukan.')
        input('--> Tekan enter untuk keluar')
        return

    required_files_dap = [
        '__init__.py',
        'A1_A_HandleDateOrigin.py',
        'A1_CleaningNSplitting.py',
        'A2_Cosmetics.py',
        'A2_CosmeticsB.py',
        'A3_PivotData.py',
        'addon_extract_excelSheet.py',
        'addon_Hapus_Baris.py',
        'helper.conf'
    ]

    required_files_dap = [
        'Hapus_Baris.py',
        'Extract_ExcelSheet.py',
        'KPI_Acc.py'
    ]

    missing_files = []
    for f in required_files_dap:
        if not os.path.exists(os.path.join(dapur_folder, f)):
            missing_files.append(f)

    if missing_files:
        print('--> Gagal: File berikut tidak ditemukan di folder Dapur:')
        for mf in missing_files:
            print(f'--> - {mf}')
        input('--> Tekan enter untuk keluar')
        return

    print('--> Membersihkan folder Dapur...')
    for file_path in glob.glob(os.path.join(dapur_folder, '*.xlsx')):
        try:
            os.remove(file_path)
        except OSError:
            pass

    print('--> Menyalin file sumber ke Dapur...')
    shutil.copy(source_file, os.path.join(dapur_folder, source_file))

    print('--> Pilihan Proses:')
    print('--> 1. Olah data full dengan python')
    print('--> 2. Cukup bersihkan data dan split data')
    
    choice = input('--> Masukkan pilihan (1/2): ')

    scripts_to_run = []
    output_target = ''
    final_name = ''

    if choice == '1':
        scripts_to_run = [
            'A1_A_HandleDateOrigin.py',
            'A1_CleaningNSplitting.py',
            'A2_Cosmetics.py',
            'A2_CosmeticsB.py',
            'A3_PivotData.py'
        ]
        output_target = 'Laporan KPI Revisi.xlsx'
        final_name = 'Laporan KPI Revisi.xlsx'
    elif choice == '2':
        scripts_to_run = [
            'A1_A_HandleDateOrigin.py',
            'A1_CleaningNSplitting.py',
            'A2_Cosmetics.py',
            'A2_CosmeticsB.py'
        ]
        output_target = 'Form Revisi_cosm_2_temp.xlsx' 
        final_name = 'Laporan Revisi Mentah.xlsx'
    else:
        print('--> Pilihan tidak valid.')
        input('--> Tekan enter untuk keluar')
        return

    print('--> Memulai proses eksekusi script...')
    for script in scripts_to_run:
        print(f'--> Menjalankan {script}...')
        try:
            subprocess.run(['python', script], cwd=dapur_folder, check=True)
        except subprocess.CalledProcessError:
            print(f'--> Gagal saat menjalankan {script}.')
            input('--> Tekan enter untuk keluar')
            return

    print('--> Memindahkan hasil akhir...')
    source_result = os.path.join(dapur_folder, output_target)
    
    if os.path.exists(source_result):
        if os.path.exists(final_name):
            os.remove(final_name)
        shutil.move(source_result, final_name)
        print(f'--> File berhasil dibuat: {final_name}')
    else:
        print(f'--> Gagal: File output {output_target} tidak ditemukan di folder Dapur.')

    print('--> Membersihkan sisa file di folder Dapur...')
    for file_path in glob.glob(os.path.join(dapur_folder, '*.xlsx')):
        try:
            os.remove(file_path)
        except OSError:
            pass
            
    input('--> Tekan enter untuk keluar')

if __name__ == '__main__':
    main()
