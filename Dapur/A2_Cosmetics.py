import pandas as pd
import openpyxl
import configparser
from fuzzywuzzy import process
import re
import os
import sys
from datetime import datetime

print("--> Memulai program perbaikan data")

input_file = "Form Revisi_temp.xlsx"
output_file = "Form Revisi_cosm_temp.xlsx"
config_file = "helper.conf"

if not os.path.exists(input_file):
    print(f"--> Gagal: File {input_file} tidak ditemukan")
    sys.exit()

if not os.path.exists(config_file):
    print(f"--> Gagal: File konfigurasi {config_file} tidak ditemukan")
    sys.exit()

print("--> Membaca konfigurasi dari helper.conf")
config = configparser.ConfigParser()
config.read(config_file)

try:
    valid_depos_str = config['MASTER_DATA']['list_depo']
    valid_names_str = config['MASTER_DATA']['list_nama_sales']
    
    MASTER_DEPO = [x.strip() for x in valid_depos_str.split(',')]
    MASTER_SALES = [x.strip() for x in valid_names_str.split(',')]
    
    print(f"--> Master Data Depo dimuat: {len(MASTER_DEPO)} item")
    print(f"--> Master Data Adm Sales dimuat: {len(MASTER_SALES)} item")
    
except KeyError as e:
    print(f"--> Gagal: Format helper.conf salah. Key tidak ditemukan: {e}")
    sys.exit()

print("--> Membaca file Excel")
df = pd.read_excel(input_file, header=0)
df.columns = df.columns.str.strip()

col_divisi = "Divisi"
col_adm = "Nama Adm Sales"
col_depo = "Depo"
col_kesalahan = "Kesalahan Admin/ Bukan"
col_tgl = "Tgl Revisi"

def clean_text_initial(text):
    if pd.isna(text): return text
    text_str = str(text).strip()
    if text_str.upper() == "NAN": return pd.NA
    return text_str.title()

def is_number(s):
    return bool(re.search(r'\d', str(s)))

def smart_date_parser(val):
    if pd.isna(val) or val == "":
        return None
    
    dt_obj = None
    if isinstance(val, datetime):
        dt_obj = val
    else:
        val_str = str(val).strip()
        
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%m/%d/%Y', '%d/%m/%y'):
            try:
                dt_obj = datetime.strptime(val_str, fmt)
                break
            except ValueError:
                continue
                
    if dt_obj:
        year = dt_obj.year
        month = dt_obj.month
        day = dt_obj.day
        
        if year < 100:
            year += 2000
        elif year < 1900: 
            year = 2000 + (year % 100)
            
        return datetime(year, month, day)
    return None

print("--> Menyeragamkan format teks awal")
cols_to_clean = [col_adm, col_depo, col_kesalahan]
for col in cols_to_clean:
    if col in df.columns:
        df[col] = df[col].apply(clean_text_initial)

if col_divisi in df.columns:
    df[col_divisi] = df[col_divisi].astype(str).str.upper()
    df[col_divisi] = df[col_divisi].replace('NAN', pd.NA)

print("--> Memperbaiki Tanggal")
if col_tgl in df.columns:
    dates_list = []
    for i in range(len(df)):
        val = df.at[i, col_tgl]
        parsed = smart_date_parser(val)
        dates_list.append(parsed)
        
    for i in range(len(dates_list)):
        if dates_list[i] is None:
            prev_d = None
            next_d = None
            for j in range(i - 1, -1, -1):
                if dates_list[j] is not None:
                    prev_d = dates_list[j]
                    break
            for j in range(i + 1, len(dates_list)):
                if dates_list[j] is not None:
                    next_d = dates_list[j]
                    break
            if prev_d: dates_list[i] = prev_d
            elif next_d: dates_list[i] = next_d
            
    df[col_tgl] = dates_list

print("--> Menjalankan Fuzzy Matching berdasarkan Master Data")

def apply_fuzzy(text, master_list, threshold=80):
    if pd.isna(text) or str(text).strip() == "":
        return text
    
    text_str = str(text)
    
    if text_str in master_list:
        return text_str
        
    best_match, score = process.extractOne(text_str, master_list)
    
    if score >= threshold:
        return best_match
    else:
        return text 

if col_depo in df.columns:
    print("--> Mengoreksi kolom Depo...")
    df[col_depo] = df[col_depo].apply(lambda x: apply_fuzzy(x, MASTER_DEPO, threshold=80))

if col_adm in df.columns:
    print("--> Mengoreksi kolom Nama Adm Sales...")
    df[col_adm] = df[col_adm].apply(lambda x: apply_fuzzy(x, MASTER_SALES, threshold=80))

print("--> Finalisasi Format Teks")
for col in cols_to_clean:
    if col in df.columns:
        df[col] = df[col].apply(clean_text_initial)

print("--> Menyimpan ke file Excel")
wb = openpyxl.Workbook()
ws = wb.active

for c_idx, col_name in enumerate(df.columns, 1):
    ws.cell(row=1, column=c_idx, value=col_name)

num_rows, num_cols = df.shape
for r in range(num_rows):
    for c in range(num_cols):
        val = df.iloc[r, c]
        if pd.isna(val): val = None
        
        cell = ws.cell(row=r + 2, column=c + 1)
        if isinstance(val, datetime):
            cell.value = val
            cell.number_format = 'DD/MM/YYYY'
        else:
            cell.value = val

print("--> Mengatur tampilan Excel")
ws.row_dimensions[1].height = 20

for col in ws.columns:
    max_len = 0
    col_letter = col[0].column_letter
    for cell in col:
        try:
            val_len = len(str(cell.value)) if cell.value else 0
            if val_len > max_len: max_len = val_len
        except: pass
    
    width = max_len + 2
    if width > 50: width = 50 
    ws.column_dimensions[col_letter].width = width

wb.save(output_file)
print(f"--> Selesai. File disimpan: {output_file}")