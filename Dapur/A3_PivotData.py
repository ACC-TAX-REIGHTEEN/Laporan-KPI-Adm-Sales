import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

input_file = 'Form Revisi_cosm_2_temp.xlsx'
output_file = 'Laporan KPI Revisi.xlsx'

print('--> Membaca file sumber...')
try:
    df = pd.read_excel(input_file)
except FileNotFoundError:
    print(f'--> Error: File {input_file} tidak ditemukan.')
    exit()

df['Tgl Revisi'] = pd.to_datetime(df['Tgl Revisi'], dayfirst=True, errors='coerce')
df = df.dropna(subset=['Tgl Revisi'])

month_map = {
    1: 'Jan', 2: 'Feb', 3: 'Mar', 4: 'Apr', 5: 'Mei', 6: 'Jun',
    7: 'Jul', 8: 'Agt', 9: 'Sep', 10: 'Okt', 11: 'Nov', 12: 'Des'
}

df['Tahun'] = df['Tgl Revisi'].dt.year
df['Bulan_Angka'] = df['Tgl Revisi'].dt.month
df['Periode'] = df['Bulan_Angka'].map(month_map) + ' ' + df['Tahun'].astype(str)
df['Periode_Sort'] = df['Tahun'].astype(str) + df['Bulan_Angka'].map('{:02d}'.format)

revisi_cols = ['AR', 'FB', 'INV', 'JC', 'JV', 'LPB', 'SJ', 'SO', 'SR', 'TFB']
for col in revisi_cols:
    df[col] = df['Revisi'].apply(lambda x: 1 if str(x).strip().upper() == col else 0)

df['Kesalahan_Admin'] = np.where(df['Kesalahan Admin/ Bukan'] == 'Kesalahan Admin', 1, 0)
df['Bukan_Kesalahan_Admin'] = np.where(df['Kesalahan Admin/ Bukan'] != 'Kesalahan Admin', 1, 0)

print('--> Memproses Laporan Depo...')
writer = pd.ExcelWriter(output_file, engine='xlsxwriter')
workbook = writer.book

unique_depos = df['Depo'].dropna().unique()
row_pointer = 0

agg_dict = {rc: 'sum' for rc in revisi_cols}
agg_dict['Kesalahan_Admin'] = 'sum'
agg_dict['Bukan_Kesalahan_Admin'] = 'sum'

final_cols_order = ['Kesalahan_Admin', 'Bukan_Kesalahan_Admin'] + revisi_cols

header_format = workbook.add_format({'bold': True, 'align': 'center', 'border': 1, 'bg_color': '#D3D3D3'})
sub_header_format = workbook.add_format({'bold': True, 'align': 'center', 'border': 1, 'bg_color': '#FFEB9C', 'font_color': '#9C5700'})

sheet_name_depo = 'Laporan Depo'
worksheet_depo = workbook.add_worksheet(sheet_name_depo)
writer.sheets[sheet_name_depo] = worksheet_depo

colors = ['#4F81BD', '#C0504D', '#9BBB59', '#8064A2', '#4BACC6', '#F79646', '#2c3e50', '#e67e22']

for idx, depo in enumerate(unique_depos):
    depo_color = colors[idx % len(colors)]
    cell_format = workbook.add_format({'bold': True, 'bg_color': depo_color, 'font_color': 'white'})
    
    worksheet_depo.write(row_pointer, 0, f"DEPO: {depo}", cell_format)
    row_pointer += 1
    
    df_depo = df[df['Depo'] == depo].copy()
    
    pivot_depo = df_depo.groupby(['Nama Adm Sales', 'Periode', 'Periode_Sort'])[list(agg_dict.keys())].sum().reset_index()
    pivot_depo = pivot_depo.sort_values(by=['Nama Adm Sales', 'Periode_Sort'])
    
    grand_total = pivot_depo[list(agg_dict.keys())].sum()
    grand_total_row = pd.DataFrame(grand_total).T
    grand_total_row['Nama Adm Sales'] = 'TOTAL BULANAN'
    grand_total_row['Periode'] = ''
    grand_total_row['Periode_Sort'] = '999999' 
    
    pivot_depo = pd.concat([pivot_depo, grand_total_row], ignore_index=True)
    pivot_depo = pivot_depo[['Nama Adm Sales', 'Periode'] + final_cols_order]
    
    header_list = pivot_depo.columns.tolist()
    for col_num, value in enumerate(header_list):
        worksheet_depo.write(row_pointer, col_num, value, header_format)
    row_pointer += 1
    
    for r_idx, row in pivot_depo.iterrows():
        for c_idx, value in enumerate(row):
            worksheet_depo.write(row_pointer, c_idx, value)
        row_pointer += 1
    
    row_pointer += 1
    worksheet_depo.write(row_pointer, 0, "RINGKASAN KUMULATIF PER ADM", sub_header_format)
    row_pointer += 1
    
    df_summary = df_depo.groupby('Nama Adm Sales')[['Kesalahan_Admin', 'Bukan_Kesalahan_Admin']].sum().reset_index()
    df_summary['TOTAL_REVISI'] = df_summary['Kesalahan_Admin'] + df_summary['Bukan_Kesalahan_Admin']
    
    summary_cols = ['Nama Adm Sales', 'Kesalahan_Admin', 'Bukan_Kesalahan_Admin', 'TOTAL_REVISI']
    
    total_summary = df_summary[['Kesalahan_Admin', 'Bukan_Kesalahan_Admin', 'TOTAL_REVISI']].sum()
    total_summary_row = pd.DataFrame(total_summary).T
    total_summary_row['Nama Adm Sales'] = 'TOTAL KUMULATIF'
    
    df_summary = pd.concat([df_summary, total_summary_row], ignore_index=True)
    
    for col_num, value in enumerate(summary_cols):
        worksheet_depo.write(row_pointer, col_num, value, header_format)
    row_pointer += 1
    
    for r_idx, row in df_summary.iterrows():
        for c_idx, col_name in enumerate(summary_cols):
            val = row[col_name]
            worksheet_depo.write(row_pointer, c_idx, val)
        row_pointer += 1
        
    row_pointer += 2

print('--> Memproses Laporan PIC...')
df_pic = df.groupby(['PIC REVISI', 'Periode_Sort', 'Periode']).size().reset_index(name='Jumlah')
pivot_pic = df_pic.pivot(index='PIC REVISI', columns=['Periode_Sort', 'Periode'], values='Jumlah').fillna(0)

pivot_pic = pivot_pic.sort_index(axis=1, level=0)
pivot_pic.columns = pivot_pic.columns.droplevel(0) 

pivot_pic['TOTAL'] = pivot_pic.sum(axis=1)

total_row_pic = pivot_pic.sum(axis=0)
total_row_pic.name = 'GRAND TOTAL'
pivot_pic = pd.concat([pivot_pic, pd.DataFrame(total_row_pic).T])

pivot_pic.to_excel(writer, sheet_name='PIC')

print('--> Menyimpan Data Asli...')
df_raw = pd.read_excel(input_file) 
df_raw.to_excel(writer, sheet_name='Data Asli', index=False)

writer.close()

print('--> Melakukan formatting akhir (Warna & Border)...')

wb = load_workbook(output_file)

ws_depo = wb['Laporan Depo']
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for row in ws_depo.iter_rows():
    for cell in row:
        if cell.value is not None:
            cell.border = thin_border

ws_pic = wb['PIC']
pic_header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
pic_total_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
pic_font = Font(color="FFFFFF", bold=True)
pic_font_black = Font(color="000000", bold=True)

for row in ws_pic.iter_rows():
    for cell in row:
        cell.border = thin_border
        if cell.row == 1:
            cell.fill = pic_header_fill
            cell.font = pic_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif cell.column == 1:
            cell.font = pic_font_black
        elif cell.value == 'GRAND TOTAL':
            for c in row:
                c.fill = pic_total_fill
                c.font = pic_font_black

for column_cells in ws_pic.columns:
    length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
    ws_pic.column_dimensions[column_cells[0].column_letter].width = length + 4

wb.save(output_file)
print(f'--> Selesai. File tersimpan di: {output_file}')