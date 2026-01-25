import openpyxl
import re
from datetime import datetime

filename = 'Form Revisi.xlsx'

print("--> Memulai proses pembersihan dan konversi data...")

wb = openpyxl.load_workbook(filename)
ws = wb.active

for row in range(2, ws.max_row + 1):
    cell_source = ws.cell(row=row, column=1)
    cell_target = ws.cell(row=row, column=5)

    cell_target.value = None

    val = cell_source.value

    if val:
        try:
            date_obj = None
            
            if isinstance(val, datetime):
                date_obj = val
            else:
                str_val = str(val).strip()
                str_val = re.sub(r'\s+', ' ', str_val)
                date_obj = datetime.strptime(str_val, "%d/%m/%Y %H:%M:%S")

            if date_obj:
                cell_target.value = date_obj
                cell_target.number_format = 'DD/MM/YYYY'
            
        except ValueError:
            print(f"--> Format tidak valid pada baris {row}")

wb.save(filename)
print("--> Selesai. Data berhasil disimpan.")