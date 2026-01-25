import openpyxl
import os
import copy

def list_excel_files():
    return [f for f in os.listdir('.') if f.endswith(('.xlsx', '.xlsm'))]

def main():
    print("=== EXCEL EXTRACTOR (FORMAT PRESERVATION) ===")
    
    files = list_excel_files()
    if not files:
        print("Tidak ditemukan file Excel (.xlsx) di folder ini.")
        return

    print("\nFile yang ditemukan:")
    for i, f in enumerate(files, 1):
        print(f"{i}. {f}")

    try:
        file_idx = int(input("\nPilih nomor file: ")) - 1
        selected_file = files[file_idx]
    except (ValueError, IndexError):
        print("Pilihan tidak valid.")
        return
        
    print(f"Membuka file asli: {selected_file}...")
    wb_original = openpyxl.load_workbook(selected_file)
    all_sheets = wb_original.sheetnames

    print("\nSheet yang tersedia:")
    print("0. EKSTRAK SEMUA (Masing-masing jadi file sendiri)")
    for i, s in enumerate(all_sheets, 1):
        print(f"{i}. {s}")

    sheet_input = input("\nPilih nomor sheet: ")
    
    if sheet_input.strip() == "0":
        selected_sheets = all_sheets
    else:
        try:
            selected_indices = [int(x.strip()) - 1 for x in sheet_input.split(',')]
            selected_sheets = [all_sheets[i] for i in selected_indices]
        except (ValueError, IndexError):
            print("Pilihan tidak valid.")
            return

    print("\nMemulai proses ekstraksi tanpa merubah format...")
    
    for target_sheet in selected_sheets:
        wb_temp = openpyxl.load_workbook(selected_file)
        
        for s_name in wb_temp.sheetnames:
            if s_name != target_sheet:
                std = wb_temp[s_name]
                wb_temp.remove(std)
                
        clean_name = "".join([c for c in target_sheet if c.isalnum() or c in (' ', '-', '_')]).strip()
        output_name = f"{clean_name}.xlsx"
        
        wb_temp.save(output_name)
        print(f" [+] Berhasil: {output_name} (Format Terjaga)")

    print("\nSEMUA PROSES SELESAI!")

if __name__ == "__main__":
    main()