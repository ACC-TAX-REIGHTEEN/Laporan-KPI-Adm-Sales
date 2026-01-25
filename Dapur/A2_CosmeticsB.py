import configparser
import re
from datetime import datetime
from copy import copy
from openpyxl import load_workbook

def load_config():
    config = configparser.ConfigParser(interpolation=None)
    config.optionxform = str
    config.read('helper.conf')
    return config

def copy_style(source_cell, target_cell):
    target_cell.font = copy(source_cell.font)
    target_cell.border = copy(source_cell.border)
    target_cell.fill = copy(source_cell.fill)
    target_cell.number_format = copy(source_cell.number_format)
    target_cell.protection = copy(source_cell.protection)
    target_cell.alignment = copy(source_cell.alignment)

def process_excel():
    config = load_config()

    input_file = 'Form Revisi_cosm_temp.xlsx'
    output_file = 'Form Revisi_cosm_2_temp.xlsx'

    print(f"--> Membaca file konfigurasi dan excel: {input_file}")
    
    try:
        wb = load_workbook(input_file)
        ws = wb.active
    except FileNotFoundError:
        print(f"--> File {input_file} tidak ditemukan.")
        return

    headers = {}
    
    for cell in ws[1]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    col_nama = headers.get("Nama Adm Sales")
    col_depo = headers.get("Depo")
    col_tgl = headers.get("Tgl Revisi")
    col_alasan = headers.get("Alasan Revisi")
    col_revisi = headers.get("Revisi")

    if not all([col_nama, col_depo, col_tgl, col_alasan, col_revisi]):
        print("--> Error: Kolom wajib tidak lengkap di file Excel.")
        return

    name_map = dict(config.items('MAPPING_NAMES'))
    depo_map = dict(config.items('MAPPING_DEPO'))
    
    fmt = config.get('SETTINGS', 'date_format')
    start_date = datetime.strptime(config.get('SETTINGS', 'start_date'), fmt)
    end_date = datetime.strptime(config.get('SETTINGS', 'end_date'), fmt)
    
    target_sales_list = [x.strip() for x in config.get('SETTINGS', 'target_sales').split(',')]
    keywords = [x.strip() for x in config.get('SETTINGS', 'keywords').split(',')]

    print("--> Membersihkan data Nama, Depo, dan Validasi Tanggal...")
    
    rows = list(ws.iter_rows(min_row=2))
    
    for row in rows:
    	
        cell_nama = row[col_nama-1]
        if cell_nama.value:
            val = str(cell_nama.value).strip()
            if val in name_map:
                cell_nama.value = name_map[val]
            elif val:
                cell_nama.value = val 
                
        cell_depo = row[col_depo-1]
        if cell_depo.value:
            val = str(cell_depo.value).strip()
            if val in depo_map:
                cell_depo.value = depo_map[val]
            elif val:
                cell_depo.value = val
                
        cell_tgl = row[col_tgl-1]
        current_idx = cell_tgl.row
        
        is_valid_date = False
        check_date = None
        
        if isinstance(cell_tgl.value, datetime):
            check_date = cell_tgl.value
        elif isinstance(cell_tgl.value, str):
            try:
                check_date = datetime.strptime(cell_tgl.value, fmt)
            except ValueError:
                pass
        
        if check_date:
            if start_date <= check_date <= end_date:
                is_valid_date = True
                
        if not is_valid_date and current_idx > 2:
            prev_cell = ws.cell(row=current_idx-1, column=col_tgl)
            cell_tgl.value = prev_cell.value
            cell_tgl.number_format = prev_cell.number_format

    print("--> Memindai baris untuk split berdasarkan keyword...")
    
    rows_to_process = []
    current_max_row = ws.max_row
    
    for r_idx in range(2, current_max_row + 1):
        name_val = ws.cell(row=r_idx, column=col_nama).value
        alasan_val = ws.cell(row=r_idx, column=col_alasan).value
        
        if name_val in target_sales_list and alasan_val:
            found_keywords = []
            alasan_str = str(alasan_val).upper()
            
            for k in keywords:
                
                pattern = r'\b' + re.escape(k) + r'\b'
                if re.search(pattern, alasan_str):
                    found_keywords.append(k)
            
            if len(found_keywords) > 0:
                rows_to_process.append((r_idx, found_keywords))

    print(f"--> Ditemukan {len(rows_to_process)} baris yang akan diproses.")
    
    rows_to_process.sort(key=lambda x: x[0], reverse=True)

    for r_idx, keys in rows_to_process:
        num_new = len(keys)
        
        ws.cell(row=r_idx, column=col_revisi).value = keys[0]
        
        if num_new > 1:
            ws.insert_rows(r_idx + 1, amount=num_new - 1)
            original_row = ws[r_idx]
            
            for i in range(1, num_new):
                new_row_idx = r_idx + i
                key = keys[i]
                
                for col_idx, cell in enumerate(original_row, 1):
                    target_cell = ws.cell(row=new_row_idx, column=col_idx)
                    target_cell.value = cell.value
                    copy_style(cell, target_cell)
                    
                ws.cell(row=new_row_idx, column=col_revisi).value = key

    print(f"--> Menyimpan hasil ke: {output_file}")
    wb.save(output_file)
    print("--> Selesai.")

if __name__ == "__main__":
    process_excel()