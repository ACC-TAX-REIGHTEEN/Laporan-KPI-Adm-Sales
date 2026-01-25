import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from copy import copy
import re

def copy_style(src_cell, dst_cell):
    """Fungsi bantu untuk menyalin format (bold, warna, border, dll)"""
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = copy(src_cell.number_format)
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)

def clean_excel_preserve_layout(input_file, output_file):
    print(f"--> Membaca file Excel: {input_file} ...")
    try:
        
        wb = openpyxl.load_workbook(input_file)
        ws = wb.active
    except Exception as e:
        print(f"--> Gagal membaca file: {e}")
        print("--> Pastikan file input adalah format .xlsx (Excel)")
        return
        
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Data Cleaned"
    
    header_map = {}
    target_col = 1
    
    for src_col in range(2, ws.max_column + 1):
        src_cell = ws.cell(row=1, column=src_col)
        header_val = src_cell.value
        
        dst_cell = out_ws.cell(row=1, column=target_col)
        dst_cell.value = header_val
        copy_style(src_cell, dst_cell)
        
        if header_val:
            header_map[str(header_val).strip()] = src_col
            
        col_letter_src = get_column_letter(src_col)
        col_letter_dst = get_column_letter(target_col)
        if col_letter_src in ws.column_dimensions:
             out_ws.column_dimensions[col_letter_dst].width = ws.column_dimensions[col_letter_src].width
             
        target_col += 1
        
    valid_keys = ["SJ", "INV", "SO", "JV", "AR", "LPB", "TFB", "FB", "JC", "SR"]
    key_map = {"SI": "INV", "Job Costing": "JC", "Retur Penjualan": "SR"}
    
    col_rev_idx = None
    col_nom_idx = None
    col_div_idx = None
    col_title_indices = []

    for name, idx in header_map.items():
        n_up = name.upper()
        if "Revisi" == name: col_rev_idx = idx
        if "Nomor Revisi" == name: col_nom_idx = idx
        if "Divisi" in n_up: col_div_idx = idx
        if any(x in n_up for x in ["Nama Adm Sales", "PIC REVISI", "Depo"]):
            col_title_indices.append(idx)

    if not col_rev_idx or not col_nom_idx:
        print("--> Peringatan: Kolom Revisi/Nomor Revisi tidak ditemukan secara tepat.")
        
    print("--> Memproses data dan memecah baris...")
    
    current_out_row = 2 
    
    for row_idx in range(2, ws.max_row + 1):
        
        row_vals = {}
        for c in range(2, ws.max_column + 1):
            row_vals[c] = ws.cell(row=row_idx, column=c).value
            
        if col_div_idx and row_vals.get(col_div_idx):
            row_vals[col_div_idx] = str(row_vals[col_div_idx]).upper()
        
        for t_idx in col_title_indices:
            if row_vals.get(t_idx):
                row_vals[t_idx] = str(row_vals[t_idx]).title()
                
        raw_rev = str(row_vals.get(col_rev_idx, "") or "").upper()
        raw_nom = str(row_vals.get(col_nom_idx, "") or "").upper()
        
        for k, v in key_map.items():
            raw_rev = raw_rev.replace(k, v)
            raw_nom = raw_nom.replace(k, v)
            
        global_keys = []
        for k in valid_keys:
            if re.search(r'\b' + k + r'\b', raw_rev):
                global_keys.append(k)
        if not global_keys and raw_rev.strip() in valid_keys:
             global_keys.append(raw_rev.strip())
             
        clean_nom = re.sub(r'[\n,;&+]', '|', raw_nom)
        chunks = [c.strip() for c in clean_nom.split('|') if c.strip()]
        
        rows_to_insert = [] 

        for chunk in chunks:
            
            numbers = re.findall(r'\d{5,}', chunk)
            if not numbers: continue
            
            local_keys = []
            for k in valid_keys:
                if re.search(r'\b' + k + r'\b', chunk):
                    local_keys.append(k)
            
            final_keys = local_keys if local_keys else global_keys
            
            if final_keys:
                for num in numbers:
                    for k in final_keys:
                        new_data = row_vals.copy()
                        new_data[col_rev_idx] = k   
                        new_data[col_nom_idx] = num 
                        rows_to_insert.append(new_data)
                        
        for data in rows_to_insert:
            out_col = 1
            
            for src_col in range(2, ws.max_column + 1):
                cell_val = data.get(src_col)
                
                new_cell = out_ws.cell(row=current_out_row, column=out_col)
                new_cell.value = cell_val
                
                orig_cell = ws.cell(row=row_idx, column=src_col)
                copy_style(orig_cell, new_cell)
                
                out_col += 1
            
            current_out_row += 1
            
    print(f"--> Menyimpan file ke: {output_file}")
    out_wb.save(output_file)
    print("--> Selesai.")

input_filename = "Form Revisi.xlsx"
output_filename = "Form Revisi_temp.xlsx"

clean_excel_preserve_layout(input_filename, output_filename)